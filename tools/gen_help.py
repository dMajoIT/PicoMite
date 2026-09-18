"""Generate the reference A:/help.txt for MMBasic's HELP command.

The topic corpus is harvested from the detailed-listing tables of
docs/PicoMite_User_Manual.docx (read only - the manual is never modified),
and emitted in the exact format cmd_help() expects.  See
docs/Help_File_Plan.md for the analysis and docs/HELP_File_Format.md for the
format contract.

    python tools/gen_help.py                 # full help.txt
    python tools/gen_help.py --short         # helpmin.txt (syntax + summary)
    python tools/gen_help.py --json out.json # the harvested corpus

Format rules enforced here (all of them are things cmd_help() requires):
  * CRLF line endings              - the seek-back in cmd_help() assumes them
  * every line < 200 characters    - the reader has a 256-byte buffer, unchecked
  * 7-bit ASCII only               - tokenise() masks input with & 0x7f
  * '~' only ever at the start of a header line
  * no trailing whitespace on a header - the match is whole-string
"""
import argparse
import glob
import os
import re
import sys
import textwrap
from collections import OrderedDict

from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
MANUAL = os.path.join(ROOT, "docs", "PicoMite_User_Manual.docx")

# docx table index -> section label.  Same tables docs/add_manual_bookmarks.py
# uses; skip is the number of header rows.
SECTIONS = [
    (14, "VARIABLE", 0),
    (15, "OPTION", 1),
    (16, "COMMAND", 0),
    (17, "FUNCTION", 0),
    (18, "OBSOLETE", 0),
]

WRAP = 78          # body text wrap column
MAXLINE = 200      # hard cap - the firmware reads into a 256-byte buffer

FOLD = {
    "\u2018": "'", "\u2019": "'", "\u201a": "'", "\u201b": "'",
    "\u201c": '"', "\u201d": '"', "\u201e": '"',
    "\u2013": "-", "\u2014": "-", "\u2010": "-", "\u2011": "-",
    "\u2026": "...", "\u2022": "-", "\u00b7": "-",
    "\u00b5": "u", "\u03bc": "u", "\u00b1": "+/-",
    "\u00b0": " deg", "\u00ba": " deg", "\u00bd": "1/2", "\u00bc": "1/4",
    "\u00d7": "x", "\u00f7": "/", "\u2264": "<=", "\u2265": ">=",
    "\u2260": "<>", "\u2192": "->", "\u2190": "<-",
    "\u00a0": " ", "\u200b": "", "\ufeff": "", "\u00ad": "",
}

# name cells that are symbols rather than keywords
SYMBOL_NAMES = {"'": "'", "\u2018": "'", "?": "?", "*": "*FILE", "/": "/*"}


# --------------------------------------------------------------------------
# docx reading
# --------------------------------------------------------------------------
def run_text(r):
    """Text of one run, with <w:tab/> rendered as spaces."""
    out = []
    for child in r:
        tag = child.tag.split("}")[1]
        if tag == "t":
            out.append(child.text or "")
        elif tag == "tab":
            out.append("    ")
        elif tag in ("br", "cr"):
            out.append(" ")
    return "".join(out)


def cell_paras(tc):
    """Paragraph texts, with Word's left indent turned into leading spaces."""
    out = []
    for p in tc.findall(qn("w:p")):
        txt = "".join(run_text(r) for r in p.findall(qn("w:r")))
        txt = txt.replace("\xa0", " ").rstrip()
        ind = p.find(qn("w:pPr") + "/" + qn("w:ind"))
        if ind is None:
            pPr = p.find(qn("w:pPr"))
            ind = pPr.find(qn("w:ind")) if pPr is not None else None
        if ind is not None and txt.strip():
            try:
                left = int(ind.get(qn("w:left")) or 0)
            except ValueError:
                left = 0
            if left > 0:
                txt = "  " + txt.lstrip()
        out.append(txt)
    return out


def blank_groups(paras):
    """Split a cell's paragraphs into blank-line separated groups."""
    groups, cur = [], []
    for t in paras:
        if t.strip():
            cur.append(t)
        elif cur:
            groups.append(cur)
            cur = []
    if cur:
        groups.append(cur)
    return groups


# --------------------------------------------------------------------------
# text handling
# --------------------------------------------------------------------------
def fold_ascii(s):
    for k, v in FOLD.items():
        s = s.replace(k, v)
    return "".join(c if 32 <= ord(c) <= 126 else " " for c in s)


def canon_name(line):
    """Canonical topic name from a syntax line, or None if it is a heading.

    'MM.INFO(EXISTS FILE fname$)'   -> 'MM.INFO(EXISTS FILE)'
    'LEFT$( string$, nbr )'         -> 'LEFT$'
    'SETPIN rx, tx, COM1'           -> 'SETPIN'
    'MATH FFT signal!(), out!()'    -> 'MATH FFT'
    'OPTION ANGLE RADIANS|DEGREES'  -> 'OPTION ANGLE'
    'Camera (OV2640)'               -> 'CAMERA(OV2640)'
    'Simple array arithmetic'       -> None   (a sub-heading, not a name)
    """
    s = fold_ascii(line).strip()
    if not s:
        return None
    if s[0] in SYMBOL_NAMES:
        # symbol commands: ' (comment), ? (print), *file (run), /* (block comment)
        if s.startswith("*"):
            return "*FILE"
        if s.startswith("/*"):
            return "/*"
        return SYMBOL_NAMES[s[0]]
    # 'OPTION TAB 2 | 3 | 4 | 8': keep the setting, drop the alternatives
    if "|" in s:
        head = s.split("|")[0].split()
        s = " ".join(head[:-1]) if len(head) > 1 else " ".join(head)
    s = re.sub(r"\s*\(\s*", "(", s, count=1)  # 'PIO (FLEVEL' -> 'PIO(FLEVEL'
    words = []
    for tok in s.split():
        if re.match(r"^[A-Z][A-Z0-9.$:]*$", tok):
            words.append(tok)
            continue
        # A bracketed keyword, e.g. MM.INFO(ADC) or Camera(OV2640).  The base may
        # be mixed case (Camera, Draw3D) but the keyword inside must be capitals,
        # which is what separates 'DEVICE(GAMEPAD ...)' from 'ABS( number )'.
        m = re.match(r"^([A-Za-z][A-Za-z0-9.$]*)\(", tok)
        if m:
            base = m.group(1)
            inner = s.split("(", 1)[1].split(")")[0]
            keys = []
            for it in inner.replace(",", " ").split():
                if re.match(r"^\.?[A-Z][A-Z0-9.$]*$", it):
                    keys.append(it)
                else:
                    break
            if keys:
                words.append("%s(%s)" % (base.upper(), " ".join(keys)))
            elif base.upper() == base:
                # ABS( number ) - a capitalised name, arguments of no interest
                words.append(base)
            # a lower-case base with no keyword inside is an argument, not a
            # name ('MATH SLICE sourcearray(), ...'): stop before it
        break
    if words:
        return " ".join(words)
    # No capitalised keyword.  A line carrying syntax punctuation is a mixed-case
    # command (TMC22xx pin, chip, ...); anything else is a layout heading.
    if not re.search(r"[(,=$%!]", s):
        return None
    return re.split(r"[=(,]", s.split()[0])[0].upper()


def common_prefix(names):
    """Longest leading word sequence shared by every name."""
    if not names:
        return ""
    split = [n.split() for n in names]
    out = []
    for parts in zip(*split):
        if len(set(parts)) != 1:
            break
        out.append(parts[0])
    return " ".join(out)


def is_preformatted(line):
    """Table rows and code samples: laid out with runs of spaces - never rewrap."""
    return bool(re.search(r"\S {3,}\S", line)) or line.startswith("    ")


def layout(lines, indent=""):
    """Fold to ASCII, wrap prose at WRAP, leave preformatted lines alone."""
    out = []
    for raw in lines:
        s = fold_ascii(raw).rstrip()
        if not s.strip():
            out.append("")
            continue
        if is_preformatted(s) and len(indent) + len(s) < MAXLINE:
            # a laid-out table row or code sample: leave the spacing alone and
            # let the firmware wrap it to the console width
            out.append(indent + s.rstrip())
            continue
        lead = re.match(r"^\s*", s).group(0)
        body = s.strip()
        wrapped = textwrap.wrap(
            body, width=WRAP - len(indent) - len(lead),
            break_long_words=False, break_on_hyphens=False,
        ) or [""]
        for w in wrapped:
            line = indent + lead + w
            while len(line) >= MAXLINE:        # a single unbreakable token
                out.append(line[:MAXLINE - 1])
                line = line[MAXLINE - 1:]
            out.append(line)
    return out


def squeeze(lines):
    """Collapse runs of blank lines and strip leading/trailing ones."""
    out = []
    for l in lines:
        if not l.strip():
            if not out or not out[-1].strip():
                continue
            out.append("")
        else:
            out.append(l.rstrip())
    while out and not out[-1].strip():
        out.pop()
    return out


def first_sentences(lines, count=2):
    """The opening of a description, for the short build."""
    head, rest = [], []
    for l in lines:
        if is_preformatted(l):
            break
        if not l.strip():
            if head:
                break
            continue
        (head if not rest else rest).append(l.strip())
    text = " ".join(head)
    parts = re.split(r"(?<=[.!?])\s+", text)
    return " ".join(parts[:count]).strip()


# --------------------------------------------------------------------------
# harvest
# --------------------------------------------------------------------------
class Topic(object):
    """One help topic: a name and one or more (syntax, description) sections.

    A topic keeps its sections separate rather than pooling all the syntax at
    the top, because the manual's wording depends on it - PRINT #nbr's entry
    opens "Same as above except ...".
    """

    def __init__(self, name, section):
        self.name = name
        self.section = section
        self.sections = []        # [(syntax lines, description lines)]
        self.stub_for = None      # name of the topic holding the description
        self.family = None        # names to point at, for a family keyword
        self.always = False       # keep the body even in the syntax-only build
        self.source = None        # supplementary manual this came from
        self.seealso = []

    @property
    def is_stub(self):
        return self.stub_for is not None

    @property
    def syntax(self):
        out = []
        for syn, _ in self.sections:
            out.extend(syn)
        return out

    @property
    def desc(self):
        out = []
        for _, d in self.sections:
            out.extend(d)
        return out

    def add(self, syn, desc):
        self.sections.append((list(syn), list(desc)))

    def extend_last(self, syn, desc):
        """Fold a layout heading into the section it belongs to."""
        if not self.sections:
            self.add(syn, desc)
            return
        s, d = self.sections[-1]
        self.sections[-1] = (s + list(syn), d + list(desc))


def harvest(path):
    doc = Document(path)
    topics = []
    stats = {"rows": 0, "aligned": 0, "merged": 0}
    for ti, section, skip in SECTIONS:
        for tr in doc.tables[ti]._tbl.findall(qn("w:tr"))[skip:]:
            tcs = tr.findall(qn("w:tc"))
            if len(tcs) < 2:
                continue
            ngroups = blank_groups(cell_paras(tcs[0]))
            dgroups = blank_groups(cell_paras(tcs[-1]))
            if not ngroups:
                continue
            stats["rows"] += 1
            if len(ngroups) == len(dgroups):
                stats["aligned"] += 1
                prev = None
                for ng, dg in zip(ngroups, dgroups):
                    name = canon_name(ng[0])
                    if not name:
                        # a layout heading ("Matrix arithmetic", "Examples:") -
                        # belongs to the entry above it, not to a topic
                        host = prev or (topics[-1] if topics else None)
                        if host is not None:
                            host.extend_last(ng, dg)
                        continue
                    t = Topic(name, section)
                    t.add(ng, dg)
                    topics.append(t)
                    prev = t
            else:
                # Several syntax forms sharing one description (or a group Word
                # laid out by eye).  Keep them together under the shared prefix
                # and make every other form reachable as a stub.
                stats["merged"] += 1
                names = [n for n in (canon_name(g[0]) for g in ngroups) if n]
                if not names:
                    continue
                primary = common_prefix(names) or names[0]
                t = Topic(primary, section)
                syn = []
                for g in ngroups:
                    syn.extend(g)
                desc = []
                for g in dgroups:
                    desc.extend(g + [""])
                t.add(syn, desc)
                topics.append(t)
                for g in ngroups:
                    nm = canon_name(g[0])
                    if not nm or nm == primary:
                        continue
                    s = Topic(nm, section)
                    s.add(g, [])
                    s.stub_for = primary
                    topics.append(s)
    return topics, stats


def merge(topics):
    """One topic per name; a real topic always beats a stub."""
    by_name = OrderedDict()
    for t in topics:
        key = t.name.upper()
        if key not in by_name:
            by_name[key] = t
            continue
        cur = by_name[key]
        if cur.is_stub and not t.is_stub:
            t.sections = cur.sections + t.sections
            by_name[key] = t
        elif t.is_stub:
            known = cur.syntax
            extra = [l for l in t.syntax if l not in known]
            if extra:
                cur.extend_last(extra, [])
        else:
            cur.sections.extend(t.sections)
    # MM.INFO(X) and MM.INFO$(X) are the same enquiry; the manual documents each
    # under one spelling, so give the other spelling a stub to land on.
    for key, t in list(by_name.items()):
        if not key.startswith("MM.INFO"):
            continue
        other = (t.name.replace("MM.INFO$(", "MM.INFO(", 1) if "MM.INFO$(" in t.name
                 else t.name.replace("MM.INFO(", "MM.INFO$(", 1))
        if other == t.name or other.upper() in by_name:
            continue
        s = Topic(other, t.section)
        s.add([other], [])
        s.stub_for = t.name
        by_name[other.upper()] = s

    # a stub pointing at a name that no longer exists is useless
    names = set(by_name)
    for key, t in list(by_name.items()):
        if t.is_stub and t.stub_for.upper() not in names:
            t.stub_for = None
    return by_name


# --------------------------------------------------------------------------
# the supplementary manuals
#
# For SPRITE, TILEMAP, RAY, FRAME, STEPPER, STRUCT, DRAW3D and the GUI controls
# the User Manual carries only a pointer row ("see the separate ..._User_Manual
# pdf"), so harvesting the manual alone leaves one topic where there should be
# forty.  These sources carry the real entries.  See docs/Help_Coverage_Review.md.
# --------------------------------------------------------------------------
NOT_A_MANUAL = (
    "bak_", "exile", "elite", "thrust", "port_", "plan", "notes", "review",
    "audit", "codemap", "manual_style", "feasibility", "briefing", "checklist",
    "demo", "armcfgen", "mmb2csub", "converter", "help_file", "release",
    "ov2640", "implementation", "picomite_user", "structures_implementation",
    # a tutorial, not a reference: every command it shows is documented in its
    # own family's manual, so it would only ever duplicate or contradict
    "game_development_guide",
)
EXTRA_SOURCES = ["generate_stepper_pdf.py", "Advanced Graphics Functions.docx"]

# The 3D manual writes the command as "3D CREATE"; the keyword is DRAW3D.
PREFIX_FIXES = [(re.compile(r"^3D\b"), "DRAW3D")]

# printed as the last line of a supplementary topic
PDF_NAMES = {
    "generate_stepper_pdf.py": "Stepper_Reference.pdf",
    "Advanced Graphics Functions.docx": "Advanced Graphics Functions.pdf",
}


def supplementary_sources():
    out = [p for p in sorted(glob.glob(os.path.join(ROOT, "docs", "*.md")))
           if not any(s in os.path.basename(p).lower() for s in NOT_A_MANUAL)]
    out += [os.path.join(ROOT, "docs", p) for p in EXTRA_SOURCES]
    return [p for p in out if os.path.exists(p)]


def pdf_name(path):
    base = os.path.basename(path)
    return PDF_NAMES.get(base, os.path.splitext(base)[0] + ".pdf")


def fix_prefix(line):
    for rx, repl in PREFIX_FIXES:
        if rx.match(line):
            return rx.sub(repl, line)
    return line


def clean_heading(text):
    t = fix_prefix(text.strip().strip("`*"))
    # "2.1 The STAR Command" - the \s+ matters, or the 3 of "3D CREATE" goes too
    t = re.sub(r"^\d+(\.\d+)*\.?\s+", "", t)
    t = re.sub(r"^The\s+", "", t)
    t = re.sub(r"\s+(Commands?|Functions?)$", "", t)
    # "SPRITE LOADPNG (RP2350 only)" - a build qualifier, not part of the name
    t = re.sub(r"\s*\([^)]*\b(only|RP2040|RP2350)\b[^)]*\)\s*$", "", t, flags=re.I)
    return fix_prefix(t)


def md_to_text(body):
    """Markdown section body -> the plain lines a console can print."""
    out = []
    for line in body.split("\n"):
        if line.strip().startswith("```"):
            continue                                  # keep the code, drop the fence
        line = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", line)   # links
        line = line.replace("**", "").replace("`", "")
        line = re.sub(r"^\s*[-*]\s+", "  - ", line)   # bullets
        if re.match(r"^\s*\|[-\s|:]+\|\s*$", line):
            continue                                  # table rule
        if line.strip().startswith("|"):
            line = "  " + "  ".join(c.strip() for c in line.strip().strip("|").split("|"))
        line = re.sub(r"^(#{1,6})\s+", "", line)      # sub-headings
        out.append(line.rstrip())
    return out


SYNTAX_LABEL = re.compile(r"^\s*(Syntax|Usage|Format)\s*:?\s*$", re.I)


def _run_of_syntax(name, lines, i):
    """Consecutive lines from i that are syntax for this topic's family."""
    out = []
    fam = name.upper().split("(")[0]
    while i < len(lines) and lines[i].strip():
        s = fix_prefix(lines[i].strip())
        got = canon_name(s)
        if not got or got.upper().split("(")[0] != fam:
            break
        out.append(s)
        i += 1
    return out, i


def split_syntax(name, lines):
    """Lift the syntax lines out of the body text.

    Most sections open with them; the SPRITE and TILEMAP manuals instead put
    them under a "Syntax:" label a paragraph or two in.
    """
    i = 0
    while i < len(lines) and not lines[i].strip():
        i += 1
    syntax, j = _run_of_syntax(name, lines, i)
    if syntax:
        return syntax, lines[j:]
    for n, line in enumerate(lines[:20]):
        if not SYNTAX_LABEL.match(line):
            continue
        syntax, j = _run_of_syntax(name, lines, n + 1)
        if syntax:
            return syntax, lines[:n] + lines[j:]
    return [], lines


TABLE_ROW = re.compile(r"^\s*\|\s*`?([^|`]+?)`?\s*\|(.*)\|?\s*$")


def md_table_entries(path):
    """The FRAME and RAY manuals list their functions as table rows:

        | `RAY(CAMX)` | Camera X position | Float |
    """
    txt = open(path, encoding="utf-8", errors="replace").read()
    for line in txt.split("\n"):
        m = TABLE_ROW.match(line)
        if not m:
            continue
        syntax = fix_prefix(m.group(1).strip())
        if "(" not in syntax:
            continue                       # commands come from the headings
        name = canon_name(syntax)
        if not name or "(" not in name:
            continue
        cells = [c.strip().strip("`") for c in m.group(2).split("|") if c.strip()]
        if not cells:
            continue
        yield name, [syntax], [cells[0]]


CODE_COMMENT = re.compile(r"^([A-Z][^']*?)\s{2,}'\s*(.+)$")


def md_code_entries(path):
    """Fenced code lines that document themselves with a trailing comment:

        FRAME(PW panel_id)   ' Returns the interior width of a panel
    """
    txt = open(path, encoding="utf-8", errors="replace").read()
    for m in re.finditer(r"```[a-z]*\n(.*?)```", txt, re.S):
        for line in m.group(1).split("\n"):
            cm = CODE_COMMENT.match(line.strip())
            if not cm:
                continue
            syntax = fix_prefix(cm.group(1).strip())
            name = canon_name(syntax)
            if name:
                yield name, [syntax], [cm.group(2).strip()]


def md_entries(path):
    """One entry per section, keyed on the heading - or, where the heading is
    prose ("Start Host (Transmitter)"), on the first line of its code block."""
    txt = open(path, encoding="utf-8", errors="replace").read()
    parts = re.split(r"(?m)^(#{2,6})[ \t]+(.*)$", txt)
    heads = [(len(parts[i]), clean_heading(parts[i + 1]),
              parts[i + 2] if i + 2 < len(parts) else "")
             for i in range(1, len(parts), 3)]
    for n, (level, head, body) in enumerate(heads):
        # A section owns its sub-sections - "Syntax", "Description", "Example" -
        # but stops at a sub-heading that is a command in its own right, which
        # gets its own topic and must not be repeated inside the parent's.
        for level2, head2, body2 in heads[n + 1:]:
            if level2 <= level or canon_name(head2):
                break
            body += "\n" + head2 + "\n" + body2
        lines = md_to_text(body)
        # The heading is usually the name.  Where it is prose ("Start Host
        # (Transmitter)") the first line of the section is, so offer both and
        # let the caller's keyword check decide which one is real.
        names = []
        head_name = canon_name(head)
        if head_name:
            names.append(head_name)
        for l in lines:
            if l.strip():
                first = canon_name(fix_prefix(l.strip()))
                if first and first not in names:
                    names.append(first)
                break
        for name in names:
            syntax, desc = split_syntax(name, lines)
            if not syntax and head != name and head_name == name:
                syntax = [head]
            yield name, syntax, desc


def pdf_script_entries(path):
    """generate_*_pdf.py: a code_block() starts an entry, multi_cell() fills it."""
    import ast

    def flush(entry):
        """A block listing a family (PEEK(STEPPER X), PEEK(STEPPER Y), ...)
        becomes one topic per name, all sharing the block's description."""
        if not entry:
            return
        names, lines, desc = entry
        for nm in names:
            own = [l for l in lines if canon_name(fix_prefix(l)) == nm]
            yield nm, (own or lines), list(desc)

    tree = ast.parse(open(path, encoding="utf-8", errors="replace").read())
    cur = None
    for node in ast.walk(tree):
        if not isinstance(node, ast.Call) or not isinstance(node.func, ast.Attribute):
            continue
        args = [a.value for a in node.args if isinstance(a, ast.Constant)
                and isinstance(a.value, str)]
        if not args:
            continue
        if node.func.attr in ("code_block", "chapter_title"):
            for e in flush(cur):
                yield e
            cur = None
            if node.func.attr == "code_block":
                lines = [l.strip() for l in args[0].split("\n") if l.strip()]
                names = []
                for l in lines:
                    nm = canon_name(fix_prefix(l))
                    if nm and nm not in names:
                        names.append(nm)
                if names:
                    cur = (names, lines, [])
        elif node.func.attr == "multi_cell" and cur:
            cur[2].extend(args[-1].split("\n"))
    for e in flush(cur):
        yield e


def docx_entries(path):
    """A syntax paragraph starts an entry; the prose under it is the body.

    Some paragraphs carry both, separated by a line break, and the reference
    tables at the end list CTRLVAL, MSGBOX and the CLICK functions.
    """
    from docx import Document
    doc = Document(path)
    cur = None
    for p in doc.paragraphs:
        block = [l.strip() for l in p.text.split("\n") if l.strip()]
        if not block:
            continue
        head = block[0]
        name = canon_name(head) if len(head) < 160 else None
        # a real entry, not a sentence that happens to start with a keyword
        if name and " " in head and not head.endswith("."):
            if cur:
                yield _drop_trailing_label(cur)
            cur = (name, [head], block[1:])
        elif cur:
            cur[2].extend(block)
    if cur:
        yield _drop_trailing_label(cur)

    for tb in doc.tables:
        for r in tb.rows:
            cells = []
            for c in r.cells:
                if not cells or c._tc is not cells[-1]._tc:
                    cells.append(c)
            if len(cells) < 2:
                continue
            syntax = " ".join(cells[0].text.split())
            name = canon_name(syntax)
            desc = [l.strip() for l in cells[-1].text.split("\n") if l.strip()]
            if name and desc:
                yield name, [syntax], desc


def _drop_trailing_label(entry):
    """The heading of the NEXT control ("Switch", "Radio Button") sits at the
    end of this one's prose - it is a label, not a sentence."""
    name, syntax, desc = entry
    while desc and len(desc[-1]) < 40 and not re.search(r"[.:;,)]$", desc[-1]):
        desc.pop()
    return name, syntax, desc


def harvest_supplementary(by_name, tokens_path=None, hidden=()):
    """Add every topic the supplementary manuals document. Returns a count."""
    added, weight = 0, {}
    # A name only counts if its first word is a real keyword; without this a
    # prose paragraph ("Several controls can ...") becomes a topic.
    tokens = firmware_tokens(tokens_path or os.path.join(ROOT, "AllCommands.h"))
    families = {t.split()[0] for t in tokens}
    for path in supplementary_sources():
        pdf = pdf_name(path)
        if path.endswith(".py"):
            entries = pdf_script_entries(path)
        elif path.endswith(".docx"):
            entries = docx_entries(path)
        else:
            # headings first: a table row is a one-line summary, and should not
            # displace a section that documents the same name properly
            entries = (list(md_entries(path)) + list(md_code_entries(path))
                       + list(md_table_entries(path)))
        for name, syntax, desc in entries:
            key = name.upper()
            if key in hidden or key.split("(")[0].split()[0] not in families:
                continue
            cur = by_name.get(key)
            if cur is not None and not cur.is_stub and not cur.family:
                continue                     # the User Manual documents it properly
            desc = squeeze([d.rstrip() for d in desc])
            if not syntax and not desc:
                continue
            size = sum(len(d) for d in desc)
            if weight.get(key, -1) >= size:
                continue        # already have a fuller entry for this name
            weight[key] = size
            if key in by_name and by_name[key].source:
                added -= 1      # replacing, not adding
            t = Topic(name, "FUNCTION" if "(" in name else "COMMAND")
            t.add(syntax or [name], desc + ["", "Full details: " + pdf])
            t.source = pdf
            by_name[key] = t
            added += 1
    return added


# Tables the manual uses for material that has no detailed-listing row of its
# own, but that people will certainly type HELP for.
EXTRA_TABLES = [
    ("OPERATORS", "COMMAND", "MMBasic operators, highest precedence first",
     [(3, "Arithmetic"), (4, "Bitwise shift"), (5, "Logical and comparison"),
      (6, "String")]),
    ("ESCAPE CODES", "COMMAND",
     "Escape sequences inside a string constant (OPTION ESCAPE must be set)",
     [(2, None)]),
]


def add_extra_topics(doc, by_name):
    for name, section, title, tables in EXTRA_TABLES:
        lines = [title, ""]
        for ti, heading in tables:
            if heading:
                lines.append(heading + ":")
            for r in doc.tables[ti].rows:
                cells = []
                for c in r.cells:
                    if not cells or c._tc is not cells[-1]._tc:
                        cells.append(c)
                if len(cells) < 2:
                    continue
                key = " ".join(fold_ascii(cells[0].text).split())
                val = " ".join(fold_ascii(cells[-1].text).split())
                if not key or not val:
                    continue        # the table's own header row
                pad = 24
                wrapped = textwrap.wrap(val, width=WRAP - pad,
                                        break_long_words=False,
                                        break_on_hyphens=False) or [""]
                lines.append("  %-*s%s" % (pad - 2, key, wrapped[0]))
                for cont in wrapped[1:]:
                    lines.append(" " * pad + cont)
            lines.append("")
        t = Topic(name, section)
        t.add([], lines)
        t.always = True          # a reference table, kept even in the tiny build
        by_name[name.upper()] = t


# Firmware tokens that are parts of a statement or operators rather than topics
# in their own right.  Mapped by hand because guessing gets them wrong.
KEYWORD_ALIASES = {
    "THEN": "IF", "TO": "FOR", "STEP": "FOR", "UNTIL": "DO", "WHILE": "DO",
    "AS": "DIM", "CASE": "SELECT CASE", "CASE ELSE": "SELECT CASE",
    "END SELECT": "SELECT CASE", "MIN": "MAX", "COLOR": "COLOUR",
    "ELSE IF": "ELSEIF", "END IF": "ENDIF",
    "AND": "OPERATORS", "OR": "OPERATORS", "XOR": "OPERATORS",
    "NOT": "OPERATORS", "INV": "OPERATORS", "MOD": "OPERATORS",
}

# PIO assembler mnemonics and words too generic to attach to one topic
ALIAS_SKIP = {
    "IN", "OUT", "SET", "VAR", "WAIT", "JMP", "MOV", "PULL", "PUSH", "NOP",
    "IRQ", "IRQ CLEAR", "IRQ NEXT", "IRQ NOWAIT", "IRQ PREV", "IRQ SET",
    "IRQ WAIT", "CALC", "TOPBOTTOM", "SIDE SET", "WRAP", "WRAP TARGET",
}

TOKEN_RE = re.compile(r'\{\(unsigned char \*\)"([^"]+)"\s*,\s*T_')

# tokenise() rewrites the name the user types into the token that implements it:
#     STR_REPLACE((char *)inpbuf, "LEFT$(", "SCHANGE$(E,", 2);
# Several user-facing functions share one token this way, to keep the function
# token table under the 128 a 7-bit token allows.  The right-hand names -
# BASE$, SCHANGE$, TOPBOTTOM - are implementation detail: nobody can type them
# and they must never appear as a topic.  Parsed rather than listed so that a
# collective added later is hidden automatically.
REWRITE_RE = re.compile(
    r'STR_REPLACE\(\(char \*\)inpbuf,\s*"([^"]+)"\s*,\s*"([^"]+)"')


def internal_tokens(mmbasic_c, documented, tokens):
    """Rewrite targets that no manual documents, i.e. the collective tokens.

    'documented' is the set of names harvested from the manuals, which keeps
    genuine aliases (SPRITE MEMORY -> BLIT MEMORY, MM.INFO$ -> MM.INFO) out of
    the result.
    """
    if not os.path.exists(mmbasic_c):
        return set()
    src = open(mmbasic_c, encoding="utf-8", errors="replace").read()
    m = re.search(r"void\s+(?:MIPS16\s+)?tokenise\s*\([^)]*\)\s*\{", src)
    if not m:                       # only the rewrites tokenise() itself does:
        return set()                # elsewhere STR_REPLACE handles board names
    depth, i = 1, m.end()
    while i < len(src) and depth:
        depth += (src[i] == "{") - (src[i] == "}")
        i += 1
    pairs = REWRITE_RE.findall(src[m.end():i])
    typed = {u.split("(")[0].upper().rstrip(",") for u, _ in pairs}
    hidden = set()
    for user, internal in pairs:
        if not (user[:1].isalpha() and internal[:1].isalpha()):
            continue                # ">=" / "<=" are operators, not tokens
        base = internal.split("(")[0].upper().rstrip(",")
        if base not in tokens:
            continue                # "GAME*MITE" -> "GAMEMITE": a board name
        # The token itself is internal when nobody types it: PEEK survives
        # because PEEK(BYTE is a source, SCHANGE$ does not because it is only
        # ever a target.
        if base not in typed and base not in documented:
            hidden.add(base)
        spelling = canon_name(internal.rstrip("(,"))
        if spelling and spelling.upper() not in documented:
            hidden.add(spelling.upper())
    return hidden


def firmware_tokens(path):
    """Every command/function name in AllCommands.h, union of all variants."""
    src = open(path, encoding="utf-8", errors="replace").read()
    out = set()
    for m in TOKEN_RE.finditer(src):
        n = m.group(1).strip().rstrip("(").upper()
        if re.match(r"^[A-Z][A-Z0-9.$]*( [A-Z0-9.$]+)*$", n):
            out.add(n)
    return out


def add_token_aliases(by_name, tokens, hidden=()):
    """Make firmware keywords that have no entry of their own reachable.

    Returns the tokens that could not be placed - genuine documentation gaps.
    """
    gaps = []
    for tk in sorted(tokens):
        if tk in by_name or tk in ALIAS_SKIP or tk in hidden:
            continue
        target = None
        if tk in KEYWORD_ALIASES and KEYWORD_ALIASES[tk].upper() in by_name:
            target = KEYWORD_ALIASES[tk]
        else:
            family = [t.name for t in by_name.values()
                      if t.name.upper().startswith(tk + " ")
                      or t.name.upper().startswith(tk + "(")]
            if family:
                t = Topic(tk, "COMMAND")
                t.add([tk], [])
                t.stub_for = None
                t.family = sorted(family, key=str.upper)
                by_name[tk] = t
                continue
            rx = re.compile(r"(?<![A-Za-z0-9_$.])" + re.escape(tk)
                            + r"(?![A-Za-z0-9_$])", re.I)
            cands = [t.name for t in by_name.values()
                     if not t.is_stub and any(rx.search(s) for s in t.syntax)]
            if len(cands) == 1:
                target = cands[0]
        if target:
            t = Topic(tk, "COMMAND")
            t.add([tk], [])
            t.stub_for = target
            by_name[tk] = t
        else:
            gaps.append(tk)
    return gaps


SEEALSO_RE = re.compile(
    r"\bSee(?: also)?(?: the)? ([A-Z][A-Z0-9.$]*(?: [A-Z][A-Z0-9.$]*)?)"
)


def add_seealso(by_name):
    for t in by_name.values():
        if t.is_stub:
            continue
        found = []
        for m in SEEALSO_RE.finditer(" ".join(t.desc)):
            cand = m.group(1).strip()
            for probe in (cand, cand.split()[0]):
                if probe.upper() in by_name and probe.upper() != t.name.upper():
                    if probe.upper() not in [f.upper() for f in found]:
                        found.append(probe)
                    break
        t.seealso = found[:6]


# --------------------------------------------------------------------------
# emit
# --------------------------------------------------------------------------
PREAMBLE_NAME = "HELP"

PREAMBLE = """\
MMBasic online help.  Type HELP followed by a command, function, option or
variable name - quotes are not needed:

    HELP PRINT
    HELP MATH FFT
    HELP MM.INFO(OPTION)

The name must match a topic exactly, so use the wildcards to explore:

    HELP MATH*        every topic whose name starts with MATH
    HELP *SPRITE*     every topic with SPRITE anywhere in the name
    HELP ???          every three-character name

Nothing is printed when no topic matches - try a wildcard.  Long entries
pause with PRESS ANY KEY.

Indexes of every topic in this file:

    HELP COMMANDS     HELP FUNCTIONS    HELP OPTIONS
    HELP VARIABLES    HELP OBSOLETE

This file is generated from the PicoMite User Manual, which carries the full
description of every entry along with the tutorial chapters.\
"""

PREAMBLE_TINY = """\
MMBasic syntax reference.  Type HELP followed by a command, function, option
or variable name - quotes are not needed:

    HELP PRINT
    HELP MATH FFT
    HELP MM.INFO(OPTION)

This file carries the syntax of every entry and nothing else.  The PicoMite
User Manual describes what each one does.

The name must match a topic exactly, so use the wildcards to explore:

    HELP MATH*        every topic whose name starts with MATH
    HELP *SPRITE*     every topic with SPRITE anywhere in the name

Nothing is printed when no topic matches - try a wildcard.

Indexes of every topic in this file:

    HELP COMMANDS     HELP FUNCTIONS    HELP OPTIONS
    HELP VARIABLES    HELP OBSOLETE\
"""

# Entries no manual documents, written here.  A collective token is hidden
# unless it appears in this table - writing the topic is the statement that the
# name is one users may type.
HANDWRITTEN = {
    "BASE$": ("FUNCTION", ["BASE$(base, number [, chars])"], """\
Returns 'number' written in the number base 'base', which may be 2 to 36.
Digits above 9 use the letters A to Z, so base 36 counts 0-9 then A-Z.

'chars' is the minimum number of characters to return.  A shorter result is
padded on the left with zeros; a longer one is returned in full, so the number
is never truncated.  If 'chars' is omitted no padding is added.

  PRINT BASE$(12, 100)      ' 84    - one hundred in base 12
  PRINT BASE$(36, 1295)     ' ZZ
  PRINT BASE$(3, 42, 6)     ' 001120
  PRINT BASE$(2, 5)         ' 101

In any base but 10 the number is treated as an unsigned 64-bit value, so
BASE$(2, -1) returns sixty-four 1s - the same convention as HEX$().

BIN$(), OCT$() and HEX$() are shortcuts for BASE$(2, ...), BASE$(8, ...) and
BASE$(16, ...).  Use BASE$ directly for any other base.

See also: BIN$, OCT$, HEX$, STR$, VAL"""),
}


def add_handwritten(by_name):
    for name, (section, syntax, body) in HANDWRITTEN.items():
        if name.upper() in by_name:
            continue                  # a manual documents it now - prefer that
        t = Topic(name, section)
        t.add(syntax, body.split("\n"))
        t.always = True
        by_name[name.upper()] = t


INDEX_TOPICS = [
    ("COMMANDS", "COMMAND", "MMBasic commands"),
    ("FUNCTIONS", "FUNCTION", "MMBasic functions"),
    ("OPTIONS", "OPTION", "OPTION settings"),
    ("VARIABLES", "VARIABLE", "Predefined read-only variables"),
    ("OBSOLETE", "OBSOLETE", "Obsolete commands and functions - do not use in new programs"),
]


def index_body(title, names):
    """Multi-column name index, 78 columns wide."""
    lines = [title, ""]
    names = sorted(names, key=str.upper)
    if not names:
        return lines
    width = min(max(len(n) for n in names) + 2, 38)
    cols = max(1, WRAP // width)
    rows = (len(names) + cols - 1) // cols
    for r in range(rows):
        cells = []
        for c in range(cols):
            i = r + c * rows
            if i < len(names):
                cells.append(names[i].ljust(width))
        lines.append("".join(cells).rstrip()[:MAXLINE])
    return lines


def build_blocks(by_name, short=False, tiny=False):
    blocks = [(PREAMBLE_NAME,
               (PREAMBLE_TINY if tiny else PREAMBLE).splitlines())]
    for topic_name, section, title in INDEX_TOPICS:
        names = [t.name for t in by_name.values()
                 if t.section == section and not t.is_stub]
        names += [t.name for t in by_name.values()
                  if t.section == section and t.is_stub]
        blocks.append((topic_name, index_body(title, sorted(set(names), key=str.upper))))

    reserved = {b[0].upper() for b in blocks}
    for key, t in sorted(by_name.items()):
        if key in reserved:
            continue
        body = []
        if t.family:
            body.append("%s is the first word of several entries:" % t.name)
            body.append("")
            body.extend(index_body("", t.family)[2:])
            body.append("")
            body.append("List them all with:  HELP %s*" % t.name)
        elif t.is_stub:
            body.extend(layout(t.syntax))
            body.append("")
            body.append("Described with the rest of its family:  HELP " + t.stub_for)
        elif tiny and not t.always:
            for syn, _ in t.sections:
                body.extend(layout(syn))
        else:
            for syn, desc in t.sections:
                if short:
                    summary = first_sentences(desc)
                    desc = [summary] if summary else []
                body.extend(layout(syn))
                body.append("")
                body.extend(layout(desc))
                body.append("")
            if t.seealso:
                body.append("")
                body.extend(layout(["See also: " + ", ".join(t.seealso)]))
        blocks.append((t.name, squeeze(body)))
    return blocks


def emit(blocks, path):
    out = []
    for name, body in blocks:
        out.append("~" + name.rstrip())
        for line in body:
            line = line.rstrip()
            if line.startswith("~"):
                line = " " + line        # '~' in column 1 would end the topic
            while len(line) >= MAXLINE:  # never truncate: fold instead
                cut = line.rfind(" ", 0, MAXLINE - 1)
                cut = cut if cut > 40 else MAXLINE - 1
                out.append(line[:cut])
                line = line[cut:].lstrip()
            out.append(line)
        out.append("")
    data = "\r\n".join(out) + "\r\n"
    with open(path, "wb") as f:
        f.write(data.encode("ascii", "replace"))
    return len(data)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--out")
    ap.add_argument("--short", action="store_true",
                    help="syntax plus a one or two sentence summary only")
    ap.add_argument("--tiny", action="store_true",
                    help="syntax only, no description at all")
    ap.add_argument("--no-supplementary", action="store_true",
                    help="User Manual only; skip the supplementary manuals")
    ap.add_argument("--manual", default=MANUAL)
    ap.add_argument("--tokens", default=os.path.join(ROOT, "AllCommands.h"),
                    help="firmware token table, cross-checked for gaps")
    ap.add_argument("--json", help="also write the harvested corpus as JSON")
    args = ap.parse_args()

    topics, stats = harvest(args.manual)
    by_name = merge(topics)
    add_extra_topics(Document(args.manual), by_name)
    add_handwritten(by_name)
    documented = set(by_name)
    hidden = internal_tokens(os.path.join(ROOT, "core", "MMBasic.c"),
                             documented, firmware_tokens(args.tokens))
    for key in hidden & set(by_name):
        del by_name[key]          # a collective token is not a user-facing name
    supp = (0 if args.no_supplementary
            else harvest_supplementary(by_name, args.tokens, hidden))
    gaps = []
    if args.tokens and os.path.exists(args.tokens):
        gaps = add_token_aliases(by_name, firmware_tokens(args.tokens), hidden)
    add_seealso(by_name)

    default = "help.txt"
    if args.tiny:
        default = "helptiny.txt"
    elif args.short:
        default = "helpmin.txt"
    out = args.out or os.path.join(ROOT, "docs", default)
    blocks = build_blocks(by_name, short=args.short, tiny=args.tiny)
    size = emit(blocks, out)

    real = sum(1 for t in by_name.values() if not t.is_stub and not t.family)
    print("manual rows      : %d (%d aligned, %d merged)"
          % (stats["rows"], stats["aligned"], stats["merged"]))
    print("topics           : %d (%d described, %d cross-reference stubs)"
          % (len(blocks), real, len(by_name) - real))
    print("supplementary    : %d topics from the separate manuals" % supp)
    print("written          : %s  (%.1f KB)" % (out, size / 1024.0))
    if hidden:
        print("collective tokens hidden: %s" % "  ".join(sorted(hidden)))
    if gaps:
        print("firmware keywords with no entry in the manual (%d):" % len(gaps))
        print("    " + "  ".join(gaps))

    if args.json:
        import json
        with open(args.json, "w", encoding="ascii", errors="replace") as f:
            json.dump([{"name": t.name, "section": t.section,
                        "syntax": [fold_ascii(s) for s in t.syntax],
                        "desc": [fold_ascii(s) for s in t.desc],
                        "stub_for": t.stub_for, "seealso": t.seealso}
                       for t in by_name.values()], f, indent=1)
        print("corpus           : %s" % args.json)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
