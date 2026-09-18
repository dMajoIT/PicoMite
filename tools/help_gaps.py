"""Find commands documented in the supplementary manuals but missing from help.txt.

Several command families - SPRITE, TILEMAP, RAY, FRAME, STEPPER, STRUCT, the GUI
controls, DRAW3D - have only a pointer entry in the User Manual ("see the
separate ..._User_Manual.pdf"), so harvesting the manual alone leaves the help
file with one topic where there should be forty.  This walks the supplementary
sources instead and reports what is missing.

    python tools/help_gaps.py                 # report by family
    python tools/help_gaps.py --by-source     # report by manual
    python tools/help_gaps.py --json out.json # machine readable

Sources are the markdown manuals in docs/, the Word original for the GUI
controls, and generate_stepper_pdf.py (whose text is embedded in the script).
A name counts as missing when its first word is a real firmware keyword from
AllCommands.h but the name itself is not a topic in docs/help.txt.
"""
import argparse
import collections
import glob
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from gen_help import canon_name, firmware_tokens          # noqa: E402

# docs/*.md that are manuals rather than plans, audits or port notes
NOT_A_MANUAL = (
    "bak_", "exile", "elite", "thrust", "port_", "plan", "notes", "review",
    "audit", "codemap", "manual_style", "feasibility", "briefing", "checklist",
    "demo", "armcfgen", "mmb2csub", "converter", "help_file", "release",
    "ov2640", "implementation", "picomite_user", "structures_implementation",
)
EXTRA_SOURCES = [
    os.path.join("docs", "generate_stepper_pdf.py"),
    os.path.join("docs", "Advanced Graphics Functions.docx"),
]

# The 3D manual writes the command as "3D CREATE"; the keyword is DRAW3D.
PREFIX_FIXES = [(re.compile(r"^3D\b"), "DRAW3D")]

CAPS = re.compile(r"^[A-Z][A-Z0-9$.]*(\(|\s|$)")

# first words that are language syntax or example-program noise, not families
NOT_A_FAMILY = {
    "PRINT", "DIM", "IF", "FOR", "DO", "REM", "END", "NEXT", "LOOP", "ON", "AS",
    "THEN", "RGB", "STR$", "EXIT", "RESTORE", "UNTIL", "SET", "CLOSE", "LOAD",
    "OPEN", "INPUT", "CLS", "LET", "SUB", "FUNCTION", "CONST", "DATA", "SELECT",
    "CASE", "ERROR", "PAUSE", "TIMER", "PIN", "PORT", "LOCAL", "LIBRARY",
    "ARRAY", "LIST", "TYPE", "COLOR", "AUTOSAVE", "XMODEM", "AND", "OR", "AT",
}

# names that are an argument value or a fragment of an example, not a topic
NOT_A_TOPIC = re.compile(
    r"^(FRAMEBUFFER WRITE [A-Z]$|SPI LCD|WAIT GPIO|USB CDC|USBKEYBOARD$"
    r"|MAP\(GLOW\)|PEEK\(VARADDR\)|RTC SETTIME|BLIT COMPRESSED PEEK"
    r"|RAY CAST RAY|PLAY SAMPLE PLAY|STAR (MOON|SATURN)$|GUI FCOLOUR RGB"
    r"|OPTION CPU$|OPTION ANGLE DEGREES$|LIBRARY LOAD|RAY COLOR$"
    r"|OPTION (PROFILING|TRACECACHE|CACHE DEBUG|CACHE SUB) OFF$"
    r"|SPRITE LOADPNG\(RP2350\)|SPRITE NOSTINTERRUPT$|TILEMAP SPRITE HIT)")


def fix_prefix(line):
    for rx, repl in PREFIX_FIXES:
        if rx.match(line):
            return rx.sub(repl, line)
    return line


def from_markdown(path):
    """Headings, fenced code and inline code that look like a syntax line."""
    out = []
    txt = open(path, encoding="utf-8", errors="replace").read()
    for line in txt.split("\n"):
        h = re.match(r"^#{2,5}\s+(.*)$", line)
        if not h:
            continue
        t = h.group(1).strip().strip("`*")
        t = re.sub(r"^\d+(\.\d+)*\.?\s*", "", t)          # "2.1 The STAR Command"
        t = re.sub(r"^The\s+", "", t)
        t = re.sub(r"\s+(Commands?|Functions?)$", "", t)
        out.append(t)
    for m in re.finditer(r"```[a-z]*\n(.*?)```", txt, re.S):
        out.extend(l.strip() for l in m.group(1).split("\n"))
    for m in re.finditer(r"`([A-Z0-9][A-Z0-9$.]*(?:\([^`]*\)|[^`\n]*))`", txt):
        out.append(m.group(1).strip())
    return out


def from_pdf_script(path):
    """generate_*_pdf.py keeps its text in code_block("...") calls."""
    txt = open(path, encoding="utf-8", errors="replace").read()
    out = []
    for m in re.finditer(r'code_block\("([^"]+)"', txt):
        out.extend(m.group(1).replace("\\n", "\n").split("\n"))
    return out


def from_docx(path):
    from docx import Document
    doc = Document(path)
    out = [p.text.strip() for p in doc.paragraphs]
    for tb in doc.tables:
        for r in tb.rows:
            for c in r.cells:
                out.append(c.text.strip().split("\n")[0])
    return [t for t in out if len(t) < 120]


def sources():
    out = [p for p in sorted(glob.glob(os.path.join(ROOT, "docs", "*.md")))
           if not any(s in os.path.basename(p).lower() for s in NOT_A_MANUAL)]
    out += [os.path.join(ROOT, p) for p in EXTRA_SOURCES]
    return [p for p in out if os.path.exists(p)]


def scan(help_path):
    tokens = firmware_tokens(os.path.join(ROOT, "AllCommands.h"))
    families = {t.split()[0] for t in tokens}
    topics = {l[1:].rstrip().upper()
              for l in open(help_path, encoding="ascii").read().split("\n")
              if l.startswith("~")}
    missing = {}
    for path in sources():
        name = os.path.basename(path)
        if path.endswith(".py"):
            lines = from_pdf_script(path)
        elif path.endswith(".docx"):
            lines = from_docx(path)
        else:
            lines = from_markdown(path)
        for raw in lines:
            line = fix_prefix(raw.strip())
            if not CAPS.match(line):
                continue
            topic = canon_name(line)
            if not topic:
                continue
            head = topic.split("(")[0].split()[0]
            if head in NOT_A_FAMILY or head not in families:
                continue
            if topic.upper() in topics or NOT_A_TOPIC.match(topic):
                continue
            e = missing.setdefault(topic, {"sources": set(), "syntax": ""})
            e["sources"].add(name)
            if len(e["syntax"]) < len(line) < 110:
                e["syntax"] = line
    return missing


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--help-file", default=os.path.join(ROOT, "docs", "help.txt"))
    ap.add_argument("--by-source", action="store_true")
    ap.add_argument("--json")
    args = ap.parse_args()

    missing = scan(args.help_file)
    if args.by_source:
        bysrc = collections.defaultdict(list)
        for topic, e in missing.items():
            for s in e["sources"]:
                bysrc[s].append(topic)
        for s in sorted(bysrc):
            print("=" * 72)
            print("%s   -- %d missing" % (s, len(bysrc[s])))
            for t in sorted(bysrc[s]):
                print("   %-30s %s" % (t, missing[t]["syntax"][:76]))
    else:
        fam = collections.defaultdict(list)
        for topic in missing:
            fam[topic.split("(")[0].split()[0]].append(topic)
        print("%d topics missing, in %d families\n" % (len(missing), len(fam)))
        for f in sorted(fam, key=lambda k: (-len(fam[k]), k)):
            print("%-10s %3d" % (f, len(fam[f])))
            for t in sorted(fam[f]):
                print("       %-30s %s" % (t, missing[t]["syntax"][:70]))
    if args.json:
        with open(args.json, "w", encoding="ascii", errors="replace") as fh:
            json.dump({k: {"sources": sorted(v["sources"]), "syntax": v["syntax"]}
                       for k, v in missing.items()}, fh, indent=1)
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
